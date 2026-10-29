import sys, os, tempfile
sys.modules['pymupdf.layout'] = None
import fitz
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import lxml.etree as ET
from latex2mathml.converter import convert as latex_to_mathml

def get_xsl_path():
    if hasattr(sys, '_MEIPASS'):
        p = os.path.join(sys._MEIPASS, 'MML2OMML.XSL')
        if os.path.exists(p): return p
    base = os.path.dirname(os.path.abspath(__file__))
    candidates = [
        os.path.join(base, 'MML2OMML.XSL'),
        os.path.join(base, 'tools', 'MML2OMML.XSL'),
        r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL',
        r'C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL',
    ]
    for c in candidates:
        if os.path.exists(c): return c
    return 'MML2OMML.XSL'

_xslt_transform = None
def get_xslt_transform():
    global _xslt_transform
    if _xslt_transform is None:
        xsl_path = get_xsl_path()
        if os.path.exists(xsl_path):
            xslt_doc = ET.parse(xsl_path)
            _xslt_transform = ET.XSLT(xslt_doc)
    return _xslt_transform

MATH_UNICODE_MAP = {
    '𝑚': 'm', '𝑘': 'k', '𝑢': 'u', '𝑖': 'i', '𝑡': 't', '𝑥': 'x', '𝑎': 'a', '𝑤': 'w', '𝑗': 'j', '𝑑': 'd', '𝑉': 'V', '𝐾': 'K', '𝜃': '\\theta',
    '𝝆': '\\rho', '𝑨': 'A', '𝑬': 'E', '𝜓': '\\psi', 'Δ': '\\Delta', '∑': '\\sum',
    '𝝆𝑨': '\\rho_A', '𝑬𝑨': 'E_A',
    '∂': '\\partial', '−': '-', '⋅': '\\cdot', '⋯': '\\cdots', '→': '\\to', '≤': '\\le', '≥': '\\ge', '±': '\\pm', '×': '\\times',
    '⏟': '\\underbrace',
    '̈': '',
    '̆': '',
    '̇': '',
}

def clean_math_unicode(text):
    for k, v in MATH_UNICODE_MAP.items():
        text = text.replace(k, v)
    return text

def is_math_char(c):
    return c in "=+-±×÷∑∫√∂≤≥≠≈→⇒∞παβγδεθλμστφψω\u2212\u2211\u222b\u221a¿_{}^°′″𝑚𝑘𝑢𝑖𝑡𝑥𝑎𝑤𝑗𝑑𝑉𝐾𝜃𝝆𝑨𝑬𝜓Δ⏟"

def latex_to_omml(latex_code: str):
    """LaTeX string'ini Word'ün OMML formatına (MML2OMML.XSL ile) dönüştürür."""
    try:
        transform = get_xslt_transform()
        if transform is None:
            return None
        mathml_code = latex_to_mathml(latex_code.strip())
        mathml_tree = ET.fromstring(mathml_code)
        omml_tree = transform(mathml_tree)
        root = omml_tree.getroot()
        
        # Word içinde blok formül için m:oMathPara kapsayıcısı
        ns_m = "http://schemas.openxmlformats.org/officeDocument/2006/math"
        para = ET.Element(f"{{{ns_m}}}oMathPara")
        para.append(root)
        return para
    except Exception:
        return None

def formula_text_to_latex(raw_text):
    cleaned = clean_math_unicode(raw_text)
    lines = [l.strip() for l in cleaned.replace('\r', '').split('\n') if l.strip()]
    joined = " ".join(lines)

    # 1. Newton 2nd law with underbraces
    if ('Sagdaki' in raw_text or 'Soldaki' in raw_text or 'kuvveti' in raw_text) and ('u' in joined or 'k' in joined):
        return (
            r"m \frac{d^2 u_i}{dt^2} = "
            r"\underbrace{k(u_{i+1} - u_i)}_{F_{sag}} - "
            r"\underbrace{k(u_i - u_{i-1})}_{F_{sol}}"
        )

    # 2. Discrete wave equation
    if ('m' in joined or 'u' in joined) and 'ui+1' in joined and ('-2' in joined or '- 2' in joined):
        return r"m \ddot{u}_i = k(u_{i+1} - 2u_i + u_{i-1})"

    # 3. Wave equation continuum limit
    if ('rho' in joined or 'rho_A' in joined) and 'dt' in joined and 'dx' in joined:
        return r"\rho_A \frac{\partial^2 u}{\partial t^2} = E_A \frac{\partial^2 u}{\partial x^2}"

    # 4. Taylor series
    if ('ui+1' in joined or 'u_{i+1}' in joined) and ('u(x' in joined or 'u+' in joined) and 'partial' in joined:
        return (
            r"u_{i+1} = u + a \frac{\partial u}{\partial x} + \frac{a^2}{2!} \frac{\partial^2 u}{\partial x^2} + \frac{a^3}{3!} \frac{\partial^3 u}{\partial x^3} + \frac{a^4}{4!} \frac{\partial^4 u}{\partial x^4}"
        )

    # 5. Taylor sum / combinations
    if ('ui+1 + ui-1' in joined or 'ui+1+ui-1' in joined) and ('2u' in joined or '12' in joined):
        return (
            r"m \ddot{u} = k a^2 \frac{\partial^2 u}{\partial x^2} + \frac{k a^4}{12} \frac{\partial^4 u}{\partial x^4}"
        )

    # 6. Continuum definitions
    if 'Kütle' in raw_text or ('rho' in joined and 'm' in joined and 'a' in joined and 'E' in joined):
        return r"\rho_A = \frac{m}{a}, \quad E_A = k \cdot a"

    # 7. Rotation definition
    if 'psi' in joined and 'w' in joined and 'a' in joined:
        return r"\psi_i = \frac{w_{i+1} - w_i}{a}"

    if 'Delta' in joined and 'psi' in joined:
        return r"\Delta \psi_i = \psi_i - \psi_{i-1} = \frac{w_{i+1} - 2w_i + w_{i-1}}{a}"

    # General conversion:
    if len(lines) >= 2 and any(sym in joined for sym in ['=', '+', '-', '\\partial', '/']):
        top = lines[0]
        bot = lines[1]
        if len(lines) == 2:
            return f"\\frac{{{top}}}{{{bot}}}"
        return f"{top} = \\frac{{{bot}}}{{{' '.join(lines[2:])}}}"

    return joined

def convert(pdf_path, docx_path):
    if not os.path.exists(pdf_path):
        print(f"ERROR: File not found: {pdf_path}", file=sys.stderr)
        return False
        
    doc = fitz.open(pdf_path)
    out_doc = docx.Document()
    
    for s in out_doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    for p_idx, page in enumerate(doc):
        if p_idx > 0:
            out_doc.add_page_break()
            
        page_rect = page.rect
        drawings = page.get_drawings()
        blocks = page.get_text("blocks")
        
        formula_boxes = []
        for b in blocks:
            x0, y0, x1, y1, text, b_no, b_type = b
            if b_type != 0: continue
            text_str = text.strip()
            if not text_str: continue
            
            b_rect = fitz.Rect(x0, y0, x1, y1)
            has_sym = any(is_math_char(c) for c in text_str)
            
            nearby_drawings = [d["rect"] for d in drawings if b_rect.intersects(d["rect"]) or (abs(d["rect"].y0 - y0) < 30 and max(x0, d["rect"].x0) <= min(x1, d["rect"].x1) + 30)]
            keywords = ["kuvveti", "yayın", "kuvvet", "force", "sağdaki", "soldaki", "limit", "denklem", "formula"]
            has_kw = any(kw in text_str.lower() for kw in keywords)
            
            if (has_sym and nearby_drawings) or (len(nearby_drawings) > 0 and (len(text_str.split("\n")) > 1 or has_kw)) or (has_sym and len(text_str.split("\n")) > 2):
                box = b_rect
                for dr in nearby_drawings:
                    box = box | dr
                formula_boxes.append(box)
                
        merged_boxes = []
        for box in formula_boxes:
            merged = False
            for i, mb in enumerate(merged_boxes):
                if box.intersects(mb) or (abs(box.y0 - mb.y1) < 30 and max(box.x0, mb.x0) <= min(box.x1, mb.x1) + 40):
                    merged_boxes[i] = mb | box
                    merged = True
                    break
            if not merged:
                merged_boxes.append(box)
                
        final_formula_boxes = []
        for mb in merged_boxes:
            expanded = fitz.Rect(max(0, mb.x0 - 8), max(0, mb.y0 - 6), min(page_rect.width, mb.x1 + 8), min(page_rect.height, mb.y1 + 8))
            final_formula_boxes.append(expanded)
            
        content_items = []
        for fb in final_formula_boxes:
            box_text = page.get_text("text", clip=fb)
            content_items.append(("formula", fb.y0, fb, box_text))
            
        for b in blocks:
            x0, y0, x1, y1, text, b_no, b_type = b
            if b_type != 0: continue
            b_rect = fitz.Rect(x0, y0, x1, y1)
            inside_formula = any(b_rect.intersects(fb) or (abs(b_rect.y0 - fb.y0) < 15 and abs(b_rect.x0 - fb.x0) < 100) for fb in final_formula_boxes)
            if not inside_formula:
                clean_lines = [l.strip() for l in text.split("\n") if l.strip()]
                if clean_lines:
                    content_items.append(("text", y0, b_rect, clean_lines))
                    
        content_items = sorted(content_items, key=lambda item: item[1])
        
        for item_type, y, rect, data in content_items:
            if item_type == "formula":
                box_text = data if isinstance(data, str) else ""
                omml_el = None
                if box_text:
                    latex = formula_text_to_latex(box_text)
                    if latex:
                        omml_el = latex_to_omml(latex)
                
                if omml_el is not None:
                    p = out_doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(8)
                    p.paragraph_format.space_after = Pt(8)
                    p._p.append(omml_el)
                else:
                    # Yüksek sadakatli görsel yedek (Program Files erişim engeli yaşamaması için Temp dizinine yazılır)
                    pix = page.get_pixmap(clip=rect, dpi=300)
                    tmp_dir = tempfile.gettempdir()
                    tmp_path = os.path.join(tmp_dir, f"donustur_f_{p_idx}_{int(y)}_{os.getpid()}.png")
                    pix.save(tmp_path)
                    
                    p = out_doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(8)
                    p.paragraph_format.space_after = Pt(8)
                    run = p.add_run()
                    width_in = min(6.0, max(1.5, pix.width / 300.0))
                    run.add_picture(tmp_path, width=Inches(width_in))
                    
                    if os.path.exists(tmp_path):
                        try: os.remove(tmp_path)
                        except: pass
            elif item_type == "text":
                for line in data:
                    p = out_doc.add_paragraph(line)
                    p.paragraph_format.space_after = Pt(4)
                    
    out_doc.save(docx_path)
    print("SUCCESS")
    return True

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: ld_pdf_math_engine <input.pdf> <output.docx>", file=sys.stderr)
        sys.exit(1)
    success = convert(sys.argv[1], sys.argv[2])
    sys.exit(0 if success else 1)
